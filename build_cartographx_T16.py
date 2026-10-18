
import argparse
import math
import re
from pathlib import Path
from collections import defaultdict, Counter

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

SEEDS = list(range(42, 52))
PROTOCOLS = [
    "lt_to_lt", "pl_to_pl", "lt_to_pl_zero_shot",
    "pl_to_lt_zero_shot", "joint_to_lt", "joint_to_pl"
]
SOURCES = {
    "carto": "cartographx_v2_ten_seed_summary_v1/all_seed_protocol_results.csv",
    "classical": "benchmark_v2_ten_seed_summary_v1/all_seed_protocol_results.csv",
    "gnn": "benchmark_v2_ten_seed_gnn_baselines_v1/all_seed_candidate_protocol_results.csv",
    "ablation": "cartographx_v2_ablations_v1/ablation_validation_all_seeds.csv",
    "graph": "cartographx_v2_graph_sensitivity_v1/all_seed_graph_sensitivity_validation.csv",
    "fraction": "cartographx_v2_training_fraction_sensitivity_v1/all_seed_fraction_validation.csv",
}
METRICS = [
    "mcc", "balanced_accuracy", "f1", "roc_auc",
    "pr_auc", "tn", "fp", "fn", "tp"
]
GEO_NAMES = {
    "city_seed_metrics.csv", "tile_seed_metrics.csv",
    "country_protocol_seed_metrics.csv",
    "seed_city_confusion.csv", "seed_tile_confusion.csv",
    "seed_protocol_confusion.csv",
    "road_feature_quartile_error_strata.csv"
}
EFF_TERMS = (
    "efficien", "runtime", "profil", "resource",
    "latency", "throughput", "memory", "complexity", "timing"
)
EFF_METRIC = re.compile(
    r"(param|time|latency|memory|throughput|flop|"
    r"peak_alloc|peak_reserv|road_per|second|millisecond)",
    re.I
)

parser = argparse.ArgumentParser()
parser.add_argument("--root", default=".")
parser.add_argument("--output", default="CartoGraphX_Master_Results.xlsx")
parser.add_argument("--overwrite", action="store_true")
args = parser.parse_args()

root = Path(args.root).resolve()
output = Path(args.output).resolve()

if not root.is_dir():
    parser.error(f"Project directory not found: {root}")
if output.exists() and not args.overwrite:
    parser.error("Output exists. Use --overwrite.")

all_files = sorted(root.rglob("*.csv"))
frames, paths = {}, {}
audit = []
counts = Counter()

def locate(suffix):
    matches = [
        f for f in all_files
        if f.as_posix().lower().endswith(suffix.lower())
    ]
    if len(matches) > 1:
        print(f"WARNING: {len(matches)} matches for {suffix}")
    return matches[0] if matches else None

for name, suffix in SOURCES.items():
    f = locate(suffix)
    if f:
        try:
            frames[name] = pd.read_csv(f, low_memory=False)
            paths[name] = f
            print(f"FOUND {name}: {len(frames[name])} rows | {f}")
        except Exception as exc:
            audit.append(("READ_ERROR", name, str(f), str(exc)))
    else:
        print(f"MISSING {name}: {suffix}")

wb = Workbook()
readme = wb.active
readme.title = "README"

for record in [
    ["CARTOGRAPH-X | 16-TABLE MASTER RESULTS"],
    ["Project root", str(root)],
    ["CSV files discovered", len(all_files)],
    ["Independent seeds", "42–51"],
    ["Statistics", "Mean, sample SD, Student-t 95% CI"],
    ["Missing values", "Blank; never replaced by zero"],
    ["T03–T04", "Test performance"],
    ["T05–T07", "Validation sensitivity"],
    ["T08", "Matched-seed test MCC differences"],
    ["T09", "Efficiency, grouped by experiment configuration"],
    ["T10–T13", "Geographic and road-feature analyses"],
    ["T14", "Thresholds and OOF diagnostics"],
    ["T15", "Directional transfer differences"],
    ["T16", "Exploratory paired significance"],
    ["Caution", "Historical test splits may have been used exploratorily"],
]:
    readme.append(record)

HEADERS = [
    "Row ID", "Model / Variant", "Protocol / Group",
    "Metric", "Evaluation / Unit"
] + [f"Seed {s}" for s in SEEDS] + [
    "N", "Mean", "Sample SD",
    "95% CI lower", "95% CI upper",
    "Mean ± SD", "Source / Notes"
]

TITLES = {
    "T01": ("T01_Dataset", "Dataset and split statistics"),
    "T02": ("T02_Configuration", "Configuration and thresholds"),
    "T03": ("T03_Baseline_Test", "Baseline test performance"),
    "T04": ("T04_Carto_Test", "CartoGraph-X test performance"),
    "T05": ("T05_Ablation_Val", "Ablation validation performance"),
    "T06": ("T06_Graph_Val", "Graph sensitivity"),
    "T07": ("T07_Fraction_Val", "Training fraction sensitivity"),
    "T08": ("T08_Paired_Diffs", "Paired MCC differences"),
    "T09": ("T09_Efficiency", "Computational efficiency"),
    "T10": ("T10_Geographic", "Geographic robustness"),
    "T11": ("T11_City", "City-level performance"),
    "T12": ("T12_Tiles", "Tile-level performance"),
    "T13": ("T13_Feature_Strata", "Road-feature error strata"),
    "T14": ("T14_Threshold_OOF", "Threshold and OOF diagnostics"),
    "T15": ("T15_Asymmetry", "Transfer asymmetry"),
    "T16": ("T16_Statistical", "Paired statistical comparisons"),
}

sheets = {}
for key, (name, title) in TITLES.items():
    ws = wb.create_sheet(name)
    ws.append([title])
    ws.append(HEADERS)
    ws.freeze_panes = "F3"
    sheets[key] = ws

seen = set()

def values_by_seed(df, metric):
    if df is None or "seed" not in df or metric not in df:
        return {}
    result = {}
    for _, rec in df[["seed", metric]].dropna().iterrows():
        try:
            seed = int(rec["seed"])
            value = float(rec[metric])
            if seed not in SEEDS or not math.isfinite(value):
                continue
            if seed in result and not math.isclose(
                result[seed], value, rel_tol=1e-10, abs_tol=1e-10
            ):
                raise ValueError(
                    f"Conflicting values for seed {seed}, metric {metric}"
                )
            result[seed] = value
        except (TypeError, OverflowError):
            continue
    return result

def emit(table, model, group, metric, split, vals, source, note=""):
    key = (table, str(model), str(group), str(metric), str(split))
    if key in seen:
        audit.append(("DUPLICATE", table, str(key), str(source)))
        return
    seen.add(key)

    clean = {}
    for seed, value in vals.items():
        try:
            seed, value = int(seed), float(value)
            if seed in SEEDS and math.isfinite(value):
                clean[seed] = value
        except (ValueError, TypeError):
            pass

    ws = sheets[table]
    r = ws.max_row + 1

    row = [
        "|".join(map(str, key[1:])),
        str(model), str(group), str(metric), str(split)
    ]
    row += [clean.get(seed) for seed in SEEDS]

    row += [
        f'=COUNT(F{r}:O{r})',
        f'=IF(P{r}=0,"",AVERAGE(F{r}:O{r}))',
        f'=IF(P{r}<2,"",STDEV(F{r}:O{r}))',
        f'=IF(P{r}<2,"",Q{r}-TINV(0.05,P{r}-1)*R{r}/SQRT(P{r}))',
        f'=IF(P{r}<2,"",Q{r}+TINV(0.05,P{r}-1)*R{r}/SQRT(P{r}))',
        f'=IF(P{r}=0,"",TEXT(Q{r},"0.0000")&IF(P{r}<2,""," ± "&TEXT(R{r},"0.0000")))',
        str(source) + (" | " + note if note else "")
    ]
    ws.append(row)
    counts[table] += len(clean)
    audit.append((
        "IMPORTED" if clean else "EMPTY",
        table, str(key), str(source)
    ))

def subset(name, **filters):
    df = frames.get(name)
    if df is None:
        return None
    for column, value in filters.items():
        if column not in df:
            return df.iloc[0:0]
        df = df[df[column] == value]
    return df

def grouped(df, columns):
    if df is None or df.empty:
        return []
    columns = [c for c in columns if c in df]
    if not columns:
        return [("all", df)]
    return df.groupby(columns, dropna=False)

def label_group(columns, keys):
    if not isinstance(keys, tuple):
        keys = (keys,)
    return " | ".join(
        f"{col}={val}" for col, val in zip(columns, keys)
    )

# --------------------------------------------------
# T01–T04: dataset, configuration, baselines, metrics
# --------------------------------------------------

for protocol in PROTOCOLS:
    carto = subset("carto", protocol=protocol)

    if carto is not None and not carto.empty:
        for metric in ["train_n", "validation_n", "test_n"]:
            emit(
                "T01", "CartoGraph-X V2", protocol,
                metric, "partition nodes",
                values_by_seed(carto, metric), paths["carto"]
            )

        for metric in ["threshold", "cv_oof_threshold"]:
            if metric in carto:
                emit(
                    "T02", "CartoGraph-X V2", protocol,
                    metric, "validation / OOF",
                    values_by_seed(carto, metric), paths["carto"]
                )

        emit(
            "T03", "CartoGraph-X V2", protocol,
            "mcc", "test",
            values_by_seed(carto, "test_mcc"), paths["carto"]
        )

        for metric in METRICS:
            col = "test_" + metric
            if col in carto:
                emit(
                    "T04", "CartoGraph-X V2", protocol,
                    metric, "test",
                    values_by_seed(carto, col), paths["carto"]
                )

    for source in ["classical", "gnn"]:
        df = subset(source, protocol=protocol)
        if df is None or df.empty:
            continue

        groupcols = ["model"]
        if source == "gnn":
            groupcols.append("variant")

        for keys, group in grouped(df, groupcols):
            label = label_group(groupcols, keys)
            emit(
                "T03", label, protocol, "mcc", "test",
                values_by_seed(group, "test_mcc"), paths[source]
            )

# --------------------------------------------------
# T05–T07: all recorded validation metrics
# --------------------------------------------------

sensitivity = [
    ("ablation", "T05", ["variant", "protocol"]),
    ("graph", "T06", ["k", "radius_m", "protocol"]),
    ("fraction", "T07", ["fraction", "protocol"])
]

for source, table, columns in sensitivity:
    df = frames.get(source)
    if df is None:
        continue

    for keys, group in grouped(df, columns):
        label = label_group(columns, keys)
        for metric in METRICS:
            col = "validation_" + metric
            if col in group:
                emit(
                    table, source, label, metric, "validation",
                    values_by_seed(group, col), paths[source]
                )

# --------------------------------------------------
# T08 and T16: matched-seed comparisons
# --------------------------------------------------

for protocol in PROTOCOLS:
    base_df = subset("carto", protocol=protocol)
    if base_df is None or base_df.empty:
        continue

    base = values_by_seed(base_df, "test_mcc")

    for source in ["classical", "gnn"]:
        df = subset(source, protocol=protocol)
        if df is None or df.empty:
            continue

        cols = ["model"] + (["variant"] if source == "gnn" else [])

        for keys, group in grouped(df, cols):
            name = label_group(cols, keys)
            other = values_by_seed(group, "test_mcc")

            paired = {
                s: base[s] - other[s]
                for s in base.keys() & other.keys()
            }

            for table in ["T08", "T16"]:
                emit(
                    table, "CartoGraph-X minus " + name,
                    protocol, "paired Δ MCC", "test",
                    paired, f"{paths['carto']} + {paths[source]}"
                )

# --------------------------------------------------
# T09: efficiency CSV discovery
# --------------------------------------------------

eff_files = [
    f for f in all_files
    if any(
        term in f.name.lower() or term in f.parent.name.lower()
        for term in EFF_TERMS
    )
]

static_ws = wb.create_sheet("T09_Static")
static_ws.append([
    "Source", "Source row", "Configuration",
    "Metric", "Value", "Note"
])

unresolved_ws = wb.create_sheet("T09_Unresolved")
unresolved_ws.append([
    "Source", "Configuration", "Metric",
    "Reason", "Seed records"
])

for file in eff_files:
    try:
        df = pd.read_csv(file, low_memory=False)
    except Exception as exc:
        audit.append(("READ_ERROR", "T09", str(file), str(exc)))
        continue

    if df.empty:
        continue

    numeric = [
        c for c in df.columns
        if EFF_METRIC.search(c)
        and pd.api.types.is_numeric_dtype(df[c])
    ]
    if not numeric:
        continue

    identity = [
        c for c in [
            "training_population", "population", "model",
            "model_family", "variant", "protocol", "device",
            "epochs", "batch_size", "precision",
            "graph_type", "timing_mode", "phase"
        ] if c in df.columns
    ]

    if "seed" in df.columns:
        for keys, group in grouped(df, identity):
            label = label_group(identity, keys)

            for metric in numeric:
                try:
                    vals = values_by_seed(group, metric)
                    emit(
                        "T09", file.stem, label,
                        metric, "measured",
                        vals, file
                    )
                except ValueError as exc:
                    unresolved_ws.append([
                        str(file), label, metric,
                        str(exc), len(group)
                    ])
    else:
        for ix, rec in df.iterrows():
            label = " | ".join(
                f"{c}={rec[c]}" for c in identity
            )
            for metric in numeric:
                if pd.notna(rec[metric]):
                    static_ws.append([
                        str(file), ix + 2, label,
                        metric, float(rec[metric]),
                        "Reported measurement; not an independent run"
                    ])

# --------------------------------------------------
# T10–T13: geographic and error analysis
# --------------------------------------------------

geo_files = [
    f for f in all_files
    if f.name.lower() in GEO_NAMES
    and any(
        term in str(f.parent).lower()
        for term in ["geographic", "error"]
    )
]

geo_tables = {
    "T10": GEO_NAMES,
    "T11": {"city_seed_metrics.csv", "seed_city_confusion.csv"},
    "T12": {"tile_seed_metrics.csv", "seed_tile_confusion.csv"},
    "T13": {"road_feature_quartile_error_strata.csv"}
}

geo_metrics = [
    "mcc", "balanced_accuracy", "f1", "roc_auc", "pr_auc",
    "n", "positives", "positive_rate",
    "predicted_positive_rate", "tn", "fp", "fn", "tp",
    "fpr", "fnr", "false_positive_rate", "false_negative_rate"
]

for file in geo_files:
    try:
        df = pd.read_csv(file, low_memory=False)
    except Exception as exc:
        audit.append(("READ_ERROR", "GEO", str(file), str(exc)))
        continue

    if df.empty or "seed" not in df:
        continue

    identity = [
        c for c in [
            "model_family", "model", "protocol",
            "country", "city", "tile_id", "tile",
            "feature", "feature_bin"
        ] if c in df.columns
    ]

    metrics = [
        c for c in geo_metrics
        if c in df.columns
        and pd.api.types.is_numeric_dtype(df[c])
    ]

    for table, allowed in geo_tables.items():
        if file.name.lower() not in allowed:
            continue

        for keys, group in grouped(df, identity):
            label = label_group(identity, keys)

            for metric in metrics:
                try:
                    emit(
                        table, file.stem, label, metric,
                        "geographic audit",
                        values_by_seed(group, metric), file
                    )
                except ValueError as exc:
                    audit.append((
                        "CONFLICT", table,
                        f"{file}: {label}: {metric}", str(exc)
                    ))

# --------------------------------------------------
# T14: threshold and OOF diagnostics
# --------------------------------------------------

for source in ["carto", "classical", "gnn"]:
    df = frames.get(source)
    if df is None:
        continue

    identity = [
        c for c in [
            "model", "variant", "protocol", "training_population"
        ] if c in df.columns
    ]

    for keys, group in grouped(df, identity):
        label = label_group(identity, keys)

        for metric in [
            "threshold", "cv_oof_threshold",
            "oof_threshold", "cv_mcc", "oof_mcc",
            "cv_roc_auc", "oof_roc_auc", "cv_mcc_at_0_5"
        ]:
            if metric in group:
                emit(
                    "T14", source, label, metric,
                    "validation / training OOF",
                    values_by_seed(group, metric), paths[source]
                )

# --------------------------------------------------
# T15: directional transfer asymmetry
# --------------------------------------------------

comparisons = [
    ("LT→PL minus PL→LT", "lt_to_pl_zero_shot", "pl_to_lt_zero_shot"),
    ("LT within minus LT→PL", "lt_to_lt", "lt_to_pl_zero_shot"),
    ("PL within minus PL→LT", "pl_to_pl", "pl_to_lt_zero_shot"),
    ("Joint→LT minus LT within", "joint_to_lt", "lt_to_lt"),
    ("Joint→PL minus PL within", "joint_to_pl", "pl_to_pl")
]

for name, p1, p2 in comparisons:
    d1 = subset("carto", protocol=p1)
    d2 = subset("carto", protocol=p2)

    if d1 is None or d2 is None:
        continue

    for metric in [
        "mcc", "balanced_accuracy", "f1", "roc_auc", "pr_auc"
    ]:
        x = values_by_seed(d1, "test_" + metric)
        y = values_by_seed(d2, "test_" + metric)

        diff = {
            s: x[s] - y[s]
            for s in x.keys() & y.keys()
        }
        emit(
            "T15", name, f"{p1} minus {p2}",
            metric, "matched-seed test difference",
            diff, paths["carto"]
        )

# --------------------------------------------------
# T16: exploratory paired t-tests
# --------------------------------------------------

ws = sheets["T16"]
for col, title in [
    ("W", "Paired t-test p"),
    ("X", "Bonferroni p"),
    ("Y", "Interpretation"),
    ("Z", "Number of comparisons")
]:
    ws[f"{col}2"] = title

protocol_counts = Counter(
    ws.cell(r, 3).value
    for r in range(3, ws.max_row + 1)
)

for r in range(3, ws.max_row + 1):
    protocol = ws.cell(r, 3).value
    ws[f"Z{r}"] = protocol_counts[protocol]

    ws[f"W{r}"] = (
        f'=IF(P{r}<2,"",'
        f'IF(R{r}=0,IF(Q{r}=0,1,0),'
        f'TDIST(ABS(Q{r}/(R{r}/SQRT(P{r}))),P{r}-1,2)))'
    )
    ws[f"X{r}"] = f'=IF(W{r}="","",MIN(1,W{r}*Z{r}))'
    ws[f"Y{r}"] = (
        f'=IF(X{r}="","Insufficient paired runs",'
        f'IF(X{r}<0.05,"Adjusted p < 0.05","Not significant"))'
    )

# --------------------------------------------------
# Source inventory and import audit
# --------------------------------------------------

inventory = wb.create_sheet("CSV_Inventory")
inventory.append(["Relative path", "Bytes", "Mapped source"])

used = set(paths.values()) | set(eff_files) | set(geo_files)

for file in all_files:
    inventory.append([
        file.relative_to(root).as_posix(),
        file.stat().st_size,
        "YES" if file in used else "NO"
    ])

audit_ws = wb.create_sheet("Import_Audit")
audit_ws.append(["Status", "Table", "Detail", "Source"])

for item in audit:
    audit_ws.append([str(x) for x in item])

for table in TITLES:
    readme.append([
        table, f"{counts[table]} independent seed measurements"
    ])

# --------------------------------------------------
# Formatting and formula validation
# --------------------------------------------------

for ws in wb:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 58
    ws.column_dimensions["D"].width = 25
    ws.column_dimensions["E"].width = 26
    ws.column_dimensions["V"].width = 70

    for col in range(6, 22):
        ws.column_dimensions[get_column_letter(col)].width = 16

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(bold=True, color="FFFFFF", size=12)

    if ws in sheets.values():
        ws.auto_filter.ref = (
            f"A2:{'Z' if ws == sheets['T16'] else 'V'}{ws.max_row}"
        )

        for cell in ws[2]:
            cell.fill = PatternFill("solid", fgColor="245B86")
            cell.font = Font(bold=True, color="FFFFFF")

        for row in ws.iter_rows(min_row=3):
            for cell in row[5:20]:
                cell.number_format = "0.0000"

    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                if "@" in cell.value or "T.INV.2T" in cell.value:
                    raise AssertionError(
                        f"Invalid formula: {ws.title}!{cell.coordinate}"
                    )

wb.calculation = CalcProperties(
    calcMode="auto",
    fullCalcOnLoad=True,
    forceFullCalc=True
)

output.parent.mkdir(parents=True, exist_ok=True)
wb.save(output)

print("\nCREATED:", output)
for table in TITLES:
    print(f"  {table}: {counts[table]} seed measurements")

print("T09 unresolved:", unresolved_ws.max_row - 1)
print("T09 static:", static_ws.max_row - 1)
print("CSV files discovered:", len(all_files))
print("Open workbook in Excel to calculate formulas.")
